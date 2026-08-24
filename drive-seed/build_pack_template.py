"""Generate the management pack template (pptx) for the Drive seed.

Usage: build_pack_template.py     (writes "Management Pack Template.pptx" here)

Five slides, placeholder-driven; build_pack.py fills {{TOKENS}} and builds
the tables. Deliberately plain corporate styling — it's the client's pack,
not Codeless Ops branding.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x2A, 0x44)
GREY = RGBColor(0x6B, 0x72, 0x80)


def _title_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(2.5))
    tf = box.text_frame
    tf.text = "Caldergate Distribution Group Ltd"
    tf.paragraphs[0].font.size = Pt(34)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = NAVY
    for text, size in [("Monthly Management Pack — {{PERIOD}}", 26),
                       ("DRAFT — built by close-pack skills from reviewer-signed figures only", 14),
                       ("Fictitious data (demo estate)", 11)]:
        p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = GREY if size < 20 else NAVY


def _section(prs, title: str, body_token: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.8))
    tf = box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(24)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = NAVY
    body = s.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12), Inches(5.6))
    body.text_frame.text = body_token
    body.text_frame.paragraphs[0].font.size = Pt(12)
    return s


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _title_slide(prs)
    _section(prs, "Cash & bank — {{PERIOD}}", "{{CASH_TABLE}}")
    _section(prs, "P&L vs budget — {{PERIOD}}", "{{VARIANCE_TABLE}}")
    _section(prs, "Commentary — flagged lines", "{{COMMENTARY}}")
    _section(prs, "Close status & sign-off", "{{SIGNOFFS}}")
    out = Path(__file__).resolve().parent / "Management Pack Template.pptx"
    prs.save(out)
    print(f"wrote {out.name}: {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    main()
