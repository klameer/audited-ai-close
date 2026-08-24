"""Generate the Drive budget workbook from the canon budget dataset.

Usage: build_budget_xlsx.py            (writes Budget FY26 - by line.xlsx here)

Source: erp-api/datasets/caldergate-close/data/budget_2026.json (the same
canon the packs tie to). One tab, lines x months, GBP thousands.
"""
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

HERE = Path(__file__).resolve().parent
BUDGET = json.loads((HERE.parent / "erp-api" / "datasets" / "caldergate-close"
                     / "data" / "budget_2026.json").read_text())
MONTHS = ["Jan-26", "Feb-26", "Mar-26", "Apr-26", "May-26", "Jun-26",
          "Jul-26", "Aug-26", "Sep-26", "Oct-26", "Nov-26", "Dec-26"]


def main() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget FY26"
    ws["A1"] = "Caldergate Distribution Group Ltd — Budget FY26 (£'000, MTD by month)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Source: Budget Model (board-approved). Fictitious data (demo estate)."
    ws.append([])
    header = ["Line"] + MONTHS
    ws.append(header)
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right")
    ws["A4"].alignment = Alignment(horizontal="left")
    for line, by_month in BUDGET["lines"].items():
        ws.append([line] + [by_month.get(m) for m in MONTHS])
    ws.column_dimensions["A"].width = 38
    for col in "BCDEFGHIJKLM":
        ws.column_dimensions[col].width = 11
    for row in ws.iter_rows(min_row=5, min_col=2):
        for cell in row:
            cell.number_format = "#,##0"
    out = HERE / "Budget FY26 - by line.xlsx"
    wb.save(out)
    print(f"wrote {out.name}: {len(BUDGET['lines'])} lines x {len(MONTHS)} months")


if __name__ == "__main__":
    main()
