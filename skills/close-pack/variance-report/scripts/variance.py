"""Compute MTD variances: trial balance actuals vs the budget workbook.

Usage: variance.py tb.json budget.xlsx <period> <flag_threshold_k> outdir/

Writes outdir/variance.json and "outdir/Variance Report - <period>.xlsx".
All figures GBP thousands (the units both sources carry). Variance sign
convention: actual - budget (revenue lines: positive = ahead of budget;
cost lines are reported the same way and read in context).
Requires openpyxl.
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font


def read_budget(path: str, period: str) -> dict:
    ws = openpyxl.load_workbook(path).active
    header = None
    budgets = {}
    for row in ws.iter_rows(values_only=True):
        if row[0] == "Line":
            header = list(row)
            continue
        if header and row[0] and isinstance(row[0], str):
            month_idx = header.index(period)
            if row[month_idx] is not None:
                budgets[row[0]] = float(row[month_idx])
    if not budgets:
        raise SystemExit(f"no budget column for {period} in {path}")
    return budgets


def main(tb_path: str, budget_path: str, period: str, threshold_k: str, outdir: str) -> None:
    tb = json.load(open(tb_path))
    assert tb["period"] == period, f"trial balance is {tb['period']}, wanted {period}"
    budgets = read_budget(budget_path, period)
    threshold = float(threshold_k)

    lines = []
    for line, vals in tb["lines"].items():
        if line not in budgets:
            raise SystemExit(f"line missing from budget: {line}")
        actual, budget = float(vals["mtd_actual"]), budgets[line]
        var = round(actual - budget, 1)
        pct = round(100 * var / budget, 1) if budget else None
        lines.append({"line": line, "mtd_actual_k": actual, "mtd_budget_k": budget,
                      "variance_k": var, "variance_pct": pct,
                      "flagged": abs(var) >= threshold})
    result = {
        "report": "MTD variance vs budget", "period": period,
        "units": "GBP thousands", "flag_threshold_k": threshold,
        "line_count": len(lines),
        "flagged_count": sum(1 for l in lines if l["flagged"]),
        "total_actual_k": round(sum(l["mtd_actual_k"] for l in lines), 1),
        "total_budget_k": round(sum(l["mtd_budget_k"] for l in lines), 1),
        "total_variance_k": round(sum(l["variance_k"] for l in lines), 1),
        "lines": lines,
    }
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    json.dump(result, open(out / "variance.json", "w"), indent=1)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Variance"
    ws["A1"] = f"Caldergate Distribution Group Ltd — Variance Report {period} (£'000, MTD)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (f"Actual per trial balance; budget per Budget FY26 workbook. "
                f"Flag threshold ±{threshold:,.0f}. Fictitious data (demo estate).")
    ws.append([])
    ws.append(["Line", "Actual", "Budget", "Variance", "Var %", "Flag"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for l in result["lines"]:
        ws.append([l["line"], l["mtd_actual_k"], l["mtd_budget_k"], l["variance_k"],
                   l["variance_pct"], "FLAG" if l["flagged"] else ""])
    ws.append(["TOTAL", result["total_actual_k"], result["total_budget_k"],
               result["total_variance_k"], None, None])
    ws[f"A{ws.max_row}"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    for col in "BCDE":
        ws.column_dimensions[col].width = 11
    for row in ws.iter_rows(min_row=5, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = "#,##0"
    wb.save(out / f"Variance Report - {period}.xlsx")
    print(f"{result['flagged_count']}/{result['line_count']} lines flagged at ±{threshold:,.0f}k; "
          f"total variance {result['total_variance_k']:+,.1f}k -> {out}")


if __name__ == "__main__":
    main(*sys.argv[1:6])
