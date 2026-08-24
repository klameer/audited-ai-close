"""Fill the management pack template from the signed step outputs.

Usage: build_pack.py template.pptx step1/reconciliation.json step2/variance.json \
                     commentary.txt audit/close-log.jsonl <period> out.pptx

Every figure placed comes from the step JSONs (reviewer-signed at the gates);
the commentary text is passed through verbatim (its own numbers are vetted
by checks.py). Requires python-pptx.
"""
import json
import sys

from pptx import Presentation
from pptx.util import Inches, Pt


def _gbp(v: float) -> str:
    return f"£{v:,.2f}"


def _fill_text(slide, token: str, text: str):
    for shape in slide.shapes:
        if shape.has_text_frame and token in shape.text_frame.text:
            shape.text_frame.text = text
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(12)
            return shape
    return None


def _replace_tokens(prs, mapping: dict):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    for token, value in mapping.items():
                        if token in run.text:
                            run.text = run.text.replace(token, value)


def _add_table(slide, anchor_shape, rows: list[list], col_widths: list[float], bold_rows=()):
    left, top = anchor_shape.left, anchor_shape.top
    anchor_shape._element.getparent().remove(anchor_shape._element)
    table = slide.shapes.add_table(len(rows), len(rows[0]), left, top,
                                   Inches(sum(col_widths)), Inches(0.3 * len(rows))).table
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = "" if val is None else str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.bold = r == 0 or r in bold_rows


def main(template, rec_path, var_path, commentary_path, log_path, period, out_path):
    rec = json.load(open(rec_path))
    var = json.load(open(var_path))
    commentary = open(commentary_path, encoding="utf-8").read().strip()
    signoffs = [json.loads(l) for l in open(log_path)] if log_path != "-" else []

    prs = Presentation(template)
    slides = list(prs.slides)

    rs = rec["reconciliation_statement"]
    stale = [i for i in rec["closing_items"] if i["stale"]]
    cash_rows = [
        ["", "GBP"],
        ["Balance per bank statement", _gbp(rs["balance_per_bank_gbp"])],
        ["Less: unpresented payments", _gbp(-rs["less_unpresented_payments_gbp"])],
        ["Add: uncleared lodgements", _gbp(rs["add_uncleared_lodgements_gbp"])],
        ["Balance per cashbook", _gbp(rs["cashbook_closing_balance_gbp"])],
        [f"Reconciling items: {len(rec['closing_items'])} (all timing); "
         f"stale items: {len(stale)} — see close log for reviewer decisions", ""],
    ]
    anchor = _fill_text(slides[1], "{{CASH_TABLE}}", "{{CASH_TABLE}}")
    _add_table(slides[1], anchor, cash_rows, [5.5, 2.5], bold_rows={4})

    var_rows = [["Line", "Actual", "Budget", "Var", "Var %", ""]]
    for l in var["lines"]:
        var_rows.append([l["line"], f"{l['mtd_actual_k']:,.0f}", f"{l['mtd_budget_k']:,.0f}",
                         f"{l['variance_k']:+,.1f}",
                         "" if l["variance_pct"] is None else f"{l['variance_pct']:+.1f}%",
                         "FLAG" if l["flagged"] else ""])
    var_rows.append(["TOTAL", f"{var['total_actual_k']:,.0f}", f"{var['total_budget_k']:,.0f}",
                     f"{var['total_variance_k']:+,.1f}", "", ""])
    anchor = _fill_text(slides[2], "{{VARIANCE_TABLE}}", "{{VARIANCE_TABLE}}")
    _add_table(slides[2], anchor, var_rows, [4.2, 1.5, 1.5, 1.5, 1.3, 0.9],
               bold_rows={len(var_rows) - 1})

    _fill_text(slides[3], "{{COMMENTARY}}", commentary)

    lines = [f"{s['step']}: signed {s['signed_at_utc']} by {s['reviewer']} "
             f"(sha256 {s['result_sha256'][:12]}…)" for s in signoffs]
    lines.append(f"Units: variance figures £'000; cash figures actual GBP. "
                 f"Flag threshold ±{var['flag_threshold_k']:,.0f}k.")
    _fill_text(slides[4], "{{SIGNOFFS}}", "\n".join(lines))

    _replace_tokens(prs, {"{{PERIOD}}": period})
    prs.save(out_path)
    print(f"pack built: {out_path} ({len(slides)} slides, "
          f"{len(var['lines'])} P&L lines, {len(signoffs)} sign-off(s))")


if __name__ == "__main__":
    main(*sys.argv[1:8])
